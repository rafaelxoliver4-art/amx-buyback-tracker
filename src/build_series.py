"""Derive the weekly and monthly buyback series and write the workbook.

    buyback_mxn     = remanente_prior + program_addition - remanente_current
    shares_bought   = shares_outstanding_prior - shares_outstanding_current
    avg_price       = buyback_mxn / shares_bought
    pct_outstanding = shares_bought / shares_outstanding_current

program_addition is 0 in normal periods. When AMX tops up the programme the
remanente JUMPS UP and, without the adjustment, the period shows a negative
buyback. Such a rise is NEVER silently absorbed: config/program_additions.yaml
is consulted first, and if the date is absent the daily reports of that week
are downloaded to isolate the exact day, a proposal is written to the config
with confirmed_by_owner: false, and an ALERT is raised.
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from pathlib import Path

import yaml

sys.path.insert(0, str(Path(__file__).resolve().parent))

import fetch_reports  # noqa: E402
import parse_report  # noqa: E402
from common import (  # noqa: E402
    CONFIG_DIR,
    PoliteSession,
    RunLog,
    load_config,
    repo_path,
    resolve_date,
)

FRAME_FIELDS = [
    "date", "remanente", "buyback_mxn", "shares_outstanding", "shares_bought",
    "avg_price", "program_addition", "pct_outstanding", "source_report_date", "pdf_url",
    "month_label", "iso_week", "week_ending_sun", "no_report",
]


# --------------------------------------------------------------------------
# labels (ruling 5 and 6)
# --------------------------------------------------------------------------
def month_label(cfg: dict, d: dt.date) -> str:
    """'May-26' - the chart's x-axis label."""
    return d.strftime(cfg["series"]["month_label_format"])


def iso_week_label(cfg: dict, d: dt.date) -> str:
    """'2026-W31'."""
    y, w, _ = d.isocalendar()
    return cfg["series"]["iso_week_format"].format(year=y, week=w)


def week_ending_sunday(d: dt.date) -> dt.date:
    """The Sunday that closes d's ISO week (ISO weeks run Mon..Sun)."""
    return d + dt.timedelta(days=7 - d.isoweekday())


# --------------------------------------------------------------------------
# program additions
# --------------------------------------------------------------------------
def load_additions() -> list[dict]:
    path = CONFIG_DIR / "program_additions.yaml"
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    out = []
    for a in data.get("additions") or []:
        out.append({**a, "date": resolve_date(a["date"])})
    return sorted(out, key=lambda a: a["date"])


DEFAULT_ADDITION_NOTES = (
    "confirm against the AGM/board resolution authorising the buyback programme "
    "top-up (BMV 'Eventos Relevantes' or 'Asambleas' for AMX, or the AMX annual "
    "report); the recompras PDF never states the programme size. Then set "
    "confirmed_by_owner: true."
)


def append_addition(date: dt.date, amount: int, source: str,
                    notes: str = DEFAULT_ADDITION_NOTES) -> None:
    """Append a proposal. Never rewrites or removes an existing entry."""
    path = CONFIG_DIR / "program_additions.yaml"
    text = path.read_text(encoding="utf-8")
    entry = (f"  - date: {date.isoformat()}\n"
             f"    amount_mxn: {amount}\n"
             f"    source: \"{source}\"\n"
             f"    notes: \"{notes}\"\n"
             f"    confirmed_by_owner: false\n")
    # an empty list must become a block sequence before items can be appended
    if re.search(r"^additions:\s*\[\s*\]\s*$", text, re.M):
        text = re.sub(r"^additions:\s*\[\s*\]\s*$", "additions:", text, count=1, flags=re.M)
        path.write_text(text.rstrip("\n") + "\n" + entry, encoding="utf-8")
        return
    path.write_text(text.rstrip("\n") + "\n" + entry, encoding="utf-8")


def addition_in(additions, lo: dt.date, hi: dt.date) -> tuple[int, list[dict]]:
    hits = [a for a in additions if lo < a["date"] <= hi]
    return sum(a["amount_mxn"] for a in hits), hits


def check_duplicate_additions(additions: list[dict], log: RunLog, window_days: int = 45) -> None:
    """A re-probe can propose an addition the config already holds under a
    slightly different date. Two equal unconfirmed amounts close together are
    almost certainly the same top-up counted twice."""
    for i, a in enumerate(additions):
        for b in additions[i + 1:]:
            if (a["amount_mxn"] == b["amount_mxn"]
                    and abs((b["date"] - a["date"]).days) <= window_days
                    and not (a.get("confirmed_by_owner") and b.get("confirmed_by_owner"))):
                log.alert(
                    f"possible DUPLICATE programme addition: {a['amount_mxn']:,} MXN on "
                    f"{a['date']} and on {b['date']} - these would double-count. "
                    "Keep one entry in config/program_additions.yaml.")


# --------------------------------------------------------------------------
# ledger -> rows
# --------------------------------------------------------------------------
def ledger_rows(cfg: dict) -> list[dict]:
    serie = cfg["issuer"]["expected_series"][0]
    inv = {r["report_date"]: r for r in fetch_reports.read_inventory(cfg)}
    rows = []
    for r in parse_report.read_ledger(cfg):
        if r["serie"] != serie:
            continue
        d = dt.date.fromisoformat(r["report_date"])
        rows.append({
            "report_date": d,
            "remanente": int(r["remanente_presente"]),
            "remanente_ultimo": int(r["remanente_ultimo"]),
            "shares_outstanding": int(r["acciones_circulacion_presente"]),
            "shares_outstanding_ultimo": int(r["acciones_circulacion_ultimo"]),
            "pdf_url": r["pdf_url"] or (inv.get(r["report_date"], {}).get("pdf_url", "")),
        })
    rows.sort(key=lambda r: r["report_date"])
    return rows


def _bucket_last(rows: list[dict], keyfn) -> list[dict]:
    buckets: dict = {}
    for r in rows:                       # rows are date-sorted
        buckets[keyfn(r["report_date"])] = r
    return [buckets[k] for k in sorted(buckets)]


# --------------------------------------------------------------------------
# frame construction
# --------------------------------------------------------------------------
def build_frame(anchors: list[dict], additions: list[dict], cfg: dict,
                log: RunLog, probe: bool = True,
                used_unconfirmed: dict | None = None) -> list[dict]:
    frame: list[dict] = []
    for i, r in enumerate(anchors):
        d = r["report_date"]
        row = {
            "date": d,
            "remanente": r["remanente"],
            "shares_outstanding": r["shares_outstanding"],
            "source_report_date": d,
            "pdf_url": r["pdf_url"],
            "program_addition": 0,
            "buyback_mxn": None,
            "shares_bought": None,
            "avg_price": None,
            "pct_outstanding": None,
            "month_label": month_label(cfg, d),
            "iso_week": iso_week_label(cfg, d),
            "week_ending_sun": week_ending_sunday(d),
            "no_report": False,
            "is_anchor": i == 0,
        }
        if i == 0:
            frame.append(row)
            continue

        prev = anchors[i - 1]
        add, hits = addition_in(additions, prev["report_date"], d)
        buyback = prev["remanente"] + add - r["remanente"]

        if buyback < 0:
            add2 = _handle_unexplained_rise(prev, r, add, cfg, log, probe)
            if add2 is not None:
                additions[:] = load_additions()
                add, hits = addition_in(additions, prev["report_date"], d)
                buyback = prev["remanente"] + add - r["remanente"]
            if buyback < 0:
                log.alert(f"{d}: remanente rose by "
                          f"{r['remanente'] - prev['remanente'] - add:,} MXN with no known "
                          f"programme addition - buyback left NEGATIVE, not absorbed",
                          code="UNEXPLAINED_RISE", affected_date=d)

        # Ruling 1: an unconfirmed addition is NOT alerted per row. It is
        # collected here and reported ONCE per run by build().
        if used_unconfirmed is not None:
            for h in hits:
                if not h.get("confirmed_by_owner", False):
                    used_unconfirmed[h["date"]] = h

        shares = prev["shares_outstanding"] - r["shares_outstanding"]
        row["program_addition"] = add
        row["buyback_mxn"] = buyback
        row["shares_bought"] = shares
        row["avg_price"] = (buyback / shares) if shares else None
        row["pct_outstanding"] = (shares / r["shares_outstanding"]) if r["shares_outstanding"] else None
        frame.append(row)
    return frame


# --------------------------------------------------------------------------
# ruling 1 (2026-08-05): the display window - a VIEW, never a deletion
# --------------------------------------------------------------------------
def apply_display_filter(cfg: dict, frame: list[dict]) -> list[dict]:
    """Trim a frame for OUTPUT only.

    Never touches the ledger, the inventory or the acceptance test - those
    always see the whole history. The derivation has already run by the time
    this is called, so the first surviving row keeps the buyback measured from
    its now-hidden predecessor.
    """
    start = (cfg.get("display") or {}).get("start_year")
    if not start:
        return frame
    return [r for r in frame if r["date"].year >= int(start)]


# --------------------------------------------------------------------------
# ruling 8 (2026-08-05): the programme-REDUCTION guard
# --------------------------------------------------------------------------
def check_implied_price_band(cfg: dict, frame: list[dict], label: str, log: RunLog) -> None:
    """A cancellation that LOWERS the remanente is arithmetically identical to
    a buyback. It betrays itself as an absurd implied average price: a lot of
    cash leaves against few or no shares.

    ALERTs only. The row is never suppressed and never adjusted.
    """
    band = (cfg.get("integrity") or {}).get("implied_price_band")
    if not band:
        return
    import statistics

    seen: list[float] = []
    for r in frame:
        px, spend, shares = r.get("avg_price"), r.get("buyback_mxn"), r.get("shares_bought")

        # cash out, no shares in - divides by zero, so the band alone misses it
        if (not shares) and spend and abs(spend) >= band["zero_share_spend_alert_mxn"]:
            log.alert(f"{label} {r['date']}: {spend:,} MXN of remanente moved with ZERO "
                      "shares retired - that is not a buyback. A programme reduction or "
                      "an undeclared addition is the likely cause; the row is reported "
                      "as measured, not adjusted.",
                      code="IMPLIED_PRICE_NO_SHARES", affected_date=r["date"])
            continue
        if px is None:
            continue

        if px < band["min_mxn"] or px > band["max_mxn"]:
            log.alert(f"{label} {r['date']}: implied average price {px:,.2f} MXN is "
                      f"outside the sanity band {band['min_mxn']}-{band['max_mxn']} - "
                      "a programme reduction, a share-structure event or a data error. "
                      "Reported as measured, not adjusted.",
                      code="IMPLIED_PRICE_BAND", affected_date=r["date"])
        elif len(seen) >= band["min_periods_for_ratio"]:
            med = statistics.median(seen[-band["median_window"]:])
            if med and max(px / med, med / px) > band["max_ratio_vs_median"]:
                log.alert(f"{label} {r['date']}: implied average price {px:,.2f} MXN "
                          f"deviates from the trailing-{band['median_window']} median of "
                          f"{med:,.2f} by more than {band['max_ratio_vs_median']}x - "
                          "a programme reduction is the likeliest cause. Reported as "
                          "measured, not adjusted.",
                          code="IMPLIED_PRICE_DEVIATION", affected_date=r["date"])
        seen.append(px)


# --------------------------------------------------------------------------
# ruling 4: never a silent gap
# --------------------------------------------------------------------------
def fill_empty_weeks(cfg: dict, frame: list[dict], log: RunLog) -> list[dict]:
    """Insert an EXPLICIT row for every ISO week with no report at all.

    buyback_mxn = 0 and shares_bought = 0 (no filings means no reported
    repurchases), remanente and shares outstanding carried forward, and
    no_report = TRUE so the gap is visible rather than inferred.
    """
    if not frame:
        return frame
    out: list[dict] = []
    for i, row in enumerate(frame):
        out.append(row)
        if i + 1 >= len(frame):
            break
        # step Monday to Monday until we reach the next row's week
        cur = row["date"] - dt.timedelta(days=row["date"].isoweekday() - 1)
        nxt = frame[i + 1]["date"]
        nxt_mon = nxt - dt.timedelta(days=nxt.isoweekday() - 1)
        cur += dt.timedelta(days=7)
        while cur < nxt_mon:
            sunday = cur + dt.timedelta(days=6)
            out.append({
                "date": sunday,
                "remanente": row["remanente"],
                "shares_outstanding": row["shares_outstanding"],
                "source_report_date": None,
                "pdf_url": "",
                "program_addition": 0,
                "buyback_mxn": 0,
                "shares_bought": 0,
                "avg_price": None,
                "pct_outstanding": 0.0,
                "month_label": month_label(cfg, sunday),
                "iso_week": iso_week_label(cfg, cur),
                "week_ending_sun": sunday,
                "no_report": True,
                "is_anchor": False,
            })
            cur += dt.timedelta(days=7)
    added = len(out) - len(frame)
    if added:
        weeks = [r["iso_week"] for r in out if r["no_report"]]
        log.notice(f"{added} ISO week(s) had no BMV report at all and were written as "
                   f"explicit zero-buyback rows (no_report = TRUE): " + ", ".join(weeks),
                   code="EMPTY_WEEK")
    return out


def fill_empty_months(cfg: dict, frame: list[dict], log: RunLog) -> list[dict]:
    """The monthly equivalent. In practice never fires - BMV files ~21
    reports a month - but a silent gap is not acceptable at any granularity."""
    if not frame:
        return frame
    out: list[dict] = []
    for i, row in enumerate(frame):
        out.append(row)
        if i + 1 >= len(frame):
            break
        y, m = row["date"].year, row["date"].month
        ny, nm = frame[i + 1]["date"].year, frame[i + 1]["date"].month
        while True:
            m += 1
            if m == 13:
                y, m = y + 1, 1
            if (y, m) >= (ny, nm):
                break
            last = (dt.date(y + (m == 12), (m % 12) + 1, 1) - dt.timedelta(days=1))
            out.append({
                "date": last,
                "remanente": row["remanente"],
                "shares_outstanding": row["shares_outstanding"],
                "source_report_date": None,
                "pdf_url": "",
                "program_addition": 0,
                "buyback_mxn": 0,
                "shares_bought": 0,
                "avg_price": None,
                "pct_outstanding": 0.0,
                "month_label": month_label(cfg, last),
                "iso_week": iso_week_label(cfg, last),
                "week_ending_sun": week_ending_sunday(last),
                "no_report": True,
                "is_anchor": False,
            })
    added = len(out) - len(frame)
    if added:
        months = [r["month_label"] for r in out if r["no_report"]]
        log.notice(f"{added} calendar month(s) had no BMV report at all and were written "
                   f"as explicit zero-buyback rows: " + ", ".join(months),
                   code="EMPTY_MONTH")
    return out


def _handle_unexplained_rise(prev, cur, known_add, cfg, log, probe) -> int | None:
    """Isolate the day of an unexplained remanente rise and propose an addition."""
    rise = cur["remanente"] + known_add - prev["remanente"]
    log.alert(f"{cur['report_date']}: unexplained remanente rise of "
              f"{cur['remanente'] - prev['remanente'] - known_add:,} MXN since {prev['report_date']}")
    if not probe:
        log.warn("probing disabled (--no-network); no addition proposed")
        return None

    sc = cfg["series"]
    lo = max(prev["report_date"] + dt.timedelta(days=1),
             cur["report_date"] - dt.timedelta(days=sc["jump_probe_days_before"]))
    hi = cur["report_date"] + dt.timedelta(days=sc["jump_probe_days_after"])
    log.info(f"probing DAILY reports {lo}..{hi} to isolate the jump")

    sess = PoliteSession(cfg, log)
    fetch_reports.download_range(cfg, sess, log, lo, hi)
    parse_report.parse_all(cfg, log)

    daily = [r for r in ledger_rows(cfg) if prev["report_date"] <= r["report_date"] <= cur["report_date"]]

    # Within one report, remanente_ultimo -> remanente_presente is pure buyback.
    # So a top-up shows up in the SEAM between reports: it is the difference
    # between yesterday's "al presente" and today's "al ultimo reporte". That
    # isolates the addition EXACTLY, with no rounding.
    jump_day, jump_amt, exact = None, 0, False
    for a, b in zip(daily, daily[1:]):
        seam = b["remanente_ultimo"] - a["remanente"]
        if seam > jump_amt:
            jump_day, jump_amt, exact = b["report_date"], seam, True
    if jump_day is None:                      # seam unusable - fall back
        for a, b in zip(daily, daily[1:]):
            delta = b["remanente"] - a["remanente"]
            if delta > jump_amt:
                jump_day, jump_amt, exact = b["report_date"], delta, False

    if jump_day is None:
        log.alert("probe found no single-day remanente rise - nothing proposed")
        return None

    # An EXACT seam is a measurement - BMV's own restatement of the previous
    # report - so it is used verbatim, round or not. Rounding is only ever a
    # fallback for when the seam cannot be measured.
    #
    # Do not assume a top-up is a round INCREMENT: on 2023-04-14 AMX reset the
    # remanente to a round TOTAL of exactly 20,000,000,000, which makes the
    # increment 1,586,249,981. Rounding that to 1.5bn left a residual rise and
    # a negative buyback. The seam is right; the roundness heuristic was not.
    step = cfg["series"]["addition_round_to_mxn"]
    if exact:
        proposed = jump_amt
        how = f"exact inter-report seam of {jump_amt:,} MXN"
        after = next((r["remanente_ultimo"] for r in daily if r["report_date"] == jump_day), None)
        if after is not None and after % step == 0:
            how += f" (the remanente RESET to a round total of {after:,} MXN)"
        elif jump_amt % step:
            near = int(round(jump_amt / step) * step)
            how += (f" - note this is {jump_amt - near:+,} MXN off a round {near:,}; "
                    "BMV restatement drift of a few pesos is a known defect")
    else:
        proposed = int(round(jump_amt / step) * step)
        how = f"observed rise of {jump_amt:,} MXN rounded to the nearest {step:,} (seam unusable)"
    append_addition(jump_day, proposed,
                    f"inferred from {how} on {jump_day} (daily probe); "
                    "AGM resolution not yet confirmed")
    log.alert(f"PROPOSED programme addition {proposed:,} MXN on {jump_day} "
              f"written to config/program_additions.yaml with confirmed_by_owner: false "
              f"- OWNER MUST CONFIRM", code="ADDITION_PROPOSED", affected_date=jump_day)
    return proposed


# --------------------------------------------------------------------------
# workbook
# --------------------------------------------------------------------------
def _style_sheet(ws, style: dict, n_cols: int, n_rows: int, *, autofilter: bool = True) -> None:
    """Header row + freeze + autofilter + widths. Restrained finance style:
    bold header, one thin bottom border, NO fill anywhere."""
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    body = Font(name=style["body_font"], size=style["body_size_pt"])
    head = Font(name=style["body_font"], size=style["body_size_pt"],
                bold=style["header_bold"])
    edge = Border(bottom=Side(style=style["header_bottom_border"],
                              color=style["header_border_color"]))

    for j in range(1, n_cols + 1):
        c = ws.cell(row=1, column=j)
        c.font = head
        c.border = edge
        c.alignment = Alignment(vertical="bottom", wrap_text=False)
    for i in range(2, n_rows + 2):
        for j in range(1, n_cols + 1):
            ws.cell(row=i, column=j).font = body

    if style.get("freeze_header"):
        ws.freeze_panes = "A2"
    if autofilter and style.get("autofilter") and n_rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{n_rows + 1}"

    # widths sized to content
    lo, hi, pad = style["column_width_min"], style["column_width_max"], style["column_width_padding"]
    for j in range(1, n_cols + 1):
        widest = 0
        for i in range(1, n_rows + 2):
            v = ws.cell(row=i, column=j).value
            if v is None:
                continue
            if isinstance(v, dt.date):
                w = len(style["date_format"])
            elif isinstance(v, float):
                w = len(f"{v:,.2f}")
            elif isinstance(v, int):
                w = len(f"{v:,}")
            else:
                w = len(str(v))
            widest = max(widest, w)
        ws.column_dimensions[get_column_letter(j)].width = min(hi, max(lo, widest + pad))


def write_workbook(cfg: dict, ledger: list[dict], weekly: list[dict],
                   monthly: list[dict], log: RunLog) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wc = cfg["workbook"]
    style = wc["style"]
    wb = Workbook()

    # ruling 1: OUTPUT-only window. The ledger written to the Raw sheet below
    # is deliberately NOT filtered - Raw is the full record.
    full_weekly, full_monthly = weekly, monthly
    weekly = apply_display_filter(cfg, weekly)
    monthly = apply_display_filter(cfg, monthly)
    if len(weekly) != len(full_weekly) or len(monthly) != len(full_monthly):
        start = cfg["display"]["start_year"]
        log.notice(f"display.start_year = {start}: showing {len(weekly)} of "
                   f"{len(full_weekly)} weekly and {len(monthly)} of {len(full_monthly)} "
                   "monthly rows. This is a VIEW - the ledger, the inventory and the "
                   "acceptance test are untouched.",
                   code="DISPLAY_WINDOW")

    # ---- Raw --------------------------------------------------------------
    ws = wb.active
    ws.title = wc["sheets"]["raw"]
    ws.append(parse_report.LEDGER_FIELDS)
    numeric = {"remanente_ultimo", "remanente_presente",
               "acciones_tesoreria_ultimo", "acciones_tesoreria_presente",
               "acciones_circulacion_ultimo", "acciones_circulacion_presente"}
    for r in ledger:
        row = []
        for k in parse_report.LEDGER_FIELDS:
            v = r.get(k, "")
            if k == "report_date" and v:
                v = dt.date.fromisoformat(v)            # a real date, not a string
            elif k in numeric and str(v).strip().lstrip("-").isdigit():
                v = int(v)
            row.append(v)
        ws.append(row)
    for j, k in enumerate(parse_report.LEDGER_FIELDS, start=1):
        fmt = "#,##0" if k in numeric else (style["date_format"] if k == "report_date" else None)
        if fmt:
            for i in range(2, len(ledger) + 2):
                ws.cell(row=i, column=j).number_format = fmt
    _style_sheet(ws, style, len(parse_report.LEDGER_FIELDS), len(ledger))

    # ---- Weekly / Monthly -------------------------------------------------
    def dump(title: str, frame: list[dict], cols: list[dict]) -> None:
        from openpyxl.styles import Font as _Font
        sh = wb.create_sheet(title)
        sh.append([c["header"] for c in cols])
        for row in frame:
            out = []
            for c in cols:
                v = row.get(c["key"])
                if c["key"] == "no_report":
                    v = "TRUE" if v else "FALSE"
                out.append(v)
            sh.append(out)
        for j, c in enumerate(cols, start=1):
            for i in range(2, len(frame) + 2):
                sh.cell(row=i, column=j).number_format = c["number_format"]
        _style_sheet(sh, style, len(cols), len(frame))

        # A negative buyback means an unhandled programme addition. Make it
        # impossible to miss: red, not buried in the log.
        red = _Font(name=style["body_font"], size=style["body_size_pt"],
                    color=style["negative_buyback_font_color"])
        jb = next((i for i, c in enumerate(cols, start=1) if c["key"] == "buyback_mxn"), None)
        if jb:
            for i in range(2, len(frame) + 2):
                cell = sh.cell(row=i, column=jb)
                if isinstance(cell.value, (int, float)) and cell.value < 0:
                    cell.font = red

    dump(wc["sheets"]["weekly"], weekly, wc["weekly_columns"])
    dump(wc["sheets"]["monthly"], monthly, wc["monthly_columns"])

    # ---- YTD (also the Cycle 2 email table) -------------------------------
    yc = wc["ytd"]
    ycols = yc["columns"]
    year = max(r["date"].year for r in monthly)
    ytd = [r for r in monthly if r["date"].year == year and not r["is_anchor"]]
    # % of shares outstanding at 31-Dec of the PRIOR year. Taken from the
    # UNFILTERED frame: a display window that starts in the current year would
    # otherwise hide the very row this denominator comes from.
    prior = [r for r in full_monthly if r["date"].year < year]
    base_shares = prior[-1]["shares_outstanding"] if prior else None
    through = max(r["date"] for r in monthly if r["date"].year == year)

    sh = wb.create_sheet(wc["sheets"]["ytd"])
    sh.append([yc["title_template"].format(year=year, through=through.strftime("%d-%b-%Y"))])
    sh.cell(row=1, column=1).font = Font(name=style["body_font"],
                                         size=style["body_size_pt"] + 2, bold=True)
    sh.append([])
    header_row = 3
    sh.append([c["header"] for c in ycols])

    def _row(r):
        return {
            "month_label": r["month_label"],
            "buyback_mxn": r["buyback_mxn"],
            "shares_bought": r["shares_bought"],
            "avg_price": r["avg_price"],
            "pct_outstanding": (r["shares_bought"] / base_shares) if base_shares else None,
        }

    for r in ytd:
        sh.append([_row(r)[c["key"]] for c in ycols])

    tot_mxn = sum(r["buyback_mxn"] for r in ytd)
    tot_sh = sum(r["shares_bought"] for r in ytd)
    total = {
        "month_label": yc["total_label"],
        "buyback_mxn": tot_mxn,
        "shares_bought": tot_sh,
        # WEIGHTED - total MXN / total shares, never a mean of the monthly means
        "avg_price": (tot_mxn / tot_sh) if tot_sh else None,
        "pct_outstanding": (tot_sh / base_shares) if base_shares else None,
    }
    sh.append([total[c["key"]] for c in ycols])

    last_row = header_row + len(ytd) + 1
    from openpyxl.styles import Border, Side
    bold = Font(name=style["body_font"], size=style["body_size_pt"], bold=True)
    body = Font(name=style["body_font"], size=style["body_size_pt"])
    edge = Border(bottom=Side(style=style["header_bottom_border"],
                              color=style["header_border_color"]))
    top = Border(top=Side(style=style["header_bottom_border"],
                          color=style["header_border_color"]))
    for j, c in enumerate(ycols, start=1):
        sh.cell(row=header_row, column=j).font = bold
        sh.cell(row=header_row, column=j).border = edge
        sh.column_dimensions[get_column_letter(j)].width = max(
            style["column_width_min"], len(c["header"]) + style["column_width_padding"] + 2)
        for i in range(header_row + 1, last_row + 1):
            cell = sh.cell(row=i, column=j)
            cell.number_format = c["number_format"]
            cell.font = bold if i == last_row else body
            if i == last_row:
                cell.border = top
    sh.cell(row=last_row, column=1).number_format = "General"
    sh.freeze_panes = f"A{header_row + 1}"

    # ---- Alerts -----------------------------------------------------------
    acols = wc["alerts_columns"]
    sh = wb.create_sheet(wc["sheets"]["alerts"])
    sh.append([c["header"] for c in acols])
    for n in log.notices:
        sh.append([n.get(c["key"], "") for c in acols])
    _style_sheet(sh, style, len(acols), len(log.notices))
    for i in range(2, len(log.notices) + 2):
        sh.cell(row=i, column=4).alignment = Alignment(wrap_text=False, vertical="top")

    # ---- Chart ------------------------------------------------------------
    # The PNG is the pixel-accurate deliverable, so it is rendered FIRST and
    # in its own guard: a problem with the native Excel chart must not be able
    # to cost us the PNG, or the workbook.
    import build_chart
    ccfg = load_config("chart.yaml")
    for what, fn in (("PNG", build_chart.render_png), ("native Excel chart", build_chart.add_excel)):
        try:
            fn(cfg, ccfg, wb, monthly, log) if fn is build_chart.add_excel \
                else fn(cfg, ccfg, monthly, log)
        except Exception as exc:
            log.alert(f"{what} generation failed: {type(exc).__name__}: {exc}",
                      code="CHART_FAILED")

    path = repo_path(wc["path"])
    wb.save(path)
    log.info(f"workbook written: {path}")
    log.info(f"YTD {year}: MXN {tot_mxn:,} / {tot_sh:,} shares / avg {tot_mxn / tot_sh:.3f}")
    return path


# --------------------------------------------------------------------------
def build(cfg: dict, log: RunLog, probe: bool = True):
    rows = ledger_rows(cfg)
    if not rows:
        log.alert("ledger is empty - nothing to build, existing data left untouched",
                  code="LEDGER_EMPTY")
        return [], []
    additions = load_additions()
    check_duplicate_additions(additions, log)

    used_unconfirmed: dict = {}
    weekly = build_frame(_bucket_last(rows, lambda d: d.isocalendar()[:2]),
                         additions, cfg, log, probe, used_unconfirmed)
    monthly = build_frame(_bucket_last(rows, lambda d: (d.year, d.month)),
                          additions, cfg, log, probe, used_unconfirmed)

    fill = cfg["series"].get("fill_empty_periods") or {}
    if fill.get("weekly"):
        weekly = fill_empty_weeks(cfg, weekly, log)
    if fill.get("monthly"):
        monthly = fill_empty_months(cfg, monthly, log)

    # ruling 8: run over the FULL frames, before any display filtering, so a
    # reduction in hidden history is still caught
    check_implied_price_band(cfg, weekly, "weekly", log)
    check_implied_price_band(cfg, monthly, "monthly", log)

    # Ruling 1: ONE line per unconfirmed addition per run - an INFO notice,
    # not a repeated ALERT. It stays visible until the owner confirms it
    # against the AGM resolution.
    for d in sorted(used_unconfirmed):
        h = used_unconfirmed[d]
        note = h.get("notes") or "needs AGM resolution"
        log.notice(f"programme addition {d} MXN {h['amount_mxn']:,} unconfirmed - {note}",
                   code="ADDITION_UNCONFIRMED", affected_date=d)
    return weekly, monthly


def main() -> int:
    ap = argparse.ArgumentParser(description="Build weekly/monthly series and the workbook.")
    ap.add_argument("--no-network", action="store_true",
                    help="do not probe daily reports when an unexplained rise is found")
    args = ap.parse_args()

    cfg = load_config()
    log = RunLog(cfg)
    weekly, monthly = build(cfg, log, probe=not args.no_network)
    if not monthly:
        return 1
    write_workbook(cfg, parse_report.read_ledger(cfg), weekly, monthly, log)
    return 1 if log.alerts else 0


if __name__ == "__main__":
    raise SystemExit(main())
